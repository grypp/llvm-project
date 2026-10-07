#!/usr/bin/env python3
"""Build the tutorial from deck.json using the supplied PowerPoint's real theme.

Usage: python build_deck.py [--template path/to/reference.pptx]
Dependencies: python-pptx, Pillow. Preview uses PyMuPDF; PDF export uses LibreOffice.
All slide text, code and diagrams are native, editable PowerPoint shapes.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import html
import json
from pathlib import Path
import re

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from tutorial_layouts import render_custom


HERE = Path(__file__).resolve().parent
NAME = "llvm-mlir-dsl-tutorial-v16"
DEFAULT_TEMPLATE = HERE / "assets" / "template.pptx"
FONT_DIR = HERE / "assets" / "fonts"
MONO_DIR = FONT_DIR
C = dict(ink="000000", gray="616161", line="D5D5D5", pale="F7F7F7",
         white="FFFFFF", green="76B900", code="484848", focus="416600",
         emphasis="E4F0D0", lightgreen="F0F6E6")


def clock(seconds):
    return f"{seconds // 60:02}:{seconds % 60:02}"


def font_file(mono=False, bold=False, light=False):
    if mono:
        return MONO_DIR / ("JetBrainsMono-Bold.ttf" if bold else "JetBrainsMono-Regular.ttf")
    return FONT_DIR / f"NVIDIASans_{'Bd' if bold else 'Lt' if light else 'Rg'}.ttf"


def measure(value, size, mono=False, bold=False, light=False):
    font = ImageFont.truetype(str(font_file(mono, bold, light)), round(size * 10))
    return font.getlength(value) / 10 / 72


def wrap(value, width, size, bold=False, mono=False, light=False):
    lines = []
    for source_line in value.split("\n"):
        if not source_line:
            lines.append("")
            continue
        words = source_line.split(" ")
        current = words.pop(0)
        for word in words:
            if measure(current + " " + word, size, mono, bold, light) <= width:
                current += " " + word
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


class Deck:
    def __init__(self, template):
        self.prs = Presentation(template)
        self.logos = [deepcopy(sh._element) for sh in self.prs.slides[0].shapes
                      if sh.name in {"Logo Text", "Logo Eye"}]
        assert len(self.logos) == 2, "Reference must contain its two editable logo shapes"
        # Retain the exact masters, layouts, theme and font parts from the reference.
        for sid in list(self.prs.slides._sldIdLst):
            self.prs.part.drop_rel(sid.rId)
            self.prs.slides._sldIdLst.remove(sid)
        self.audit = []
        self.i = 0
        assert abs(self.prs.slide_width / 914400 - 10) < .001
        assert abs(self.prs.slide_height / 914400 - 5.625) < .001

    def rect(self, s, x, y, w, h, fill, stroke=None):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if stroke:
            sh.line.color.rgb = RGBColor.from_string(stroke)
            sh.line.width = Pt(.7)
        else:
            sh.line.fill.background()
        style = sh._element.find(qn("p:style"))
        if style is not None:
            sh._element.remove(style)
        sh._element.spPr.append(OxmlElement("a:effectLst"))
        return sh

    def arrow(self, s, x1, y1, x2, y2, color=None, head=True):
        sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        sh.line.color.rgb = RGBColor.from_string(color or C["focus"])
        sh.line.width = Pt(1.25)
        if head:
            tip = OxmlElement("a:tailEnd")
            tip.set("type", "triangle")
            sh._element.spPr.find(qn("a:ln")).append(tip)

    def pipeline_node(self, s, node, cx, y, label_width=1.01, accent=False):
        """Small editable document / gear / chip figures, with labels outside."""
        color = C["focus"] if accent else C["gray"]
        kind = node["kind"]
        if kind == "source":
            # Document outline, folded corner, and source-code mark.
            self.rect(s, cx - .15, y, .30, .38, C["white"], color)
            self.arrow(s, cx + .04, y, cx + .04, y + .105, color, head=False)
            self.arrow(s, cx + .04, y + .105, cx + .15, y + .105, color, head=False)
            self.arrow(s, cx + .04, y, cx + .15, y + .105, color, head=False)
            self.text(s, "</>", cx - .128, y + .16, .256, .17, 7.7, color, mono=True, align="center")
        elif kind in {"compiler", "interpreter"}:
            sh = s.shapes.add_shape(MSO_SHAPE.GEAR_6, Inches(cx - .195), Inches(y), Inches(.39), Inches(.39))
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(C["emphasis"] if accent else C["pale"])
            sh.line.color.rgb = RGBColor.from_string(color)
            sh.line.width = Pt(.9)
            style = sh._element.find(qn("p:style"))
            if style is not None:
                sh._element.remove(style)
            sh._element.spPr.append(OxmlElement("a:effectLst"))
        elif kind == "binary":
            self.rect(s, cx - .15, y + .04, .30, .30, C["lightgreen"] if accent else C["pale"], color)
            for off in (-.08, .0, .08):
                self.arrow(s, cx + off, y, cx + off, y + .04, color, head=False)
                self.arrow(s, cx + off, y + .34, cx + off, y + .38, color, head=False)
                self.arrow(s, cx - .19, y + .19 + off, cx - .15, y + .19 + off, color, head=False)
                self.arrow(s, cx + .15, y + .19 + off, cx + .19, y + .19 + off, color, head=False)
            self.text(s, "01", cx - .13, y + .12, .26, .20, 9, color, mono=True, align="center")
        else:
            raise ValueError(kind)
        self.text(s, node["label"], cx - label_width / 2, y + .45, label_width, .40,
                  10.5, C["focus"] if accent else C["ink"],
                  mono=kind == "source" or node["label"] == "mlir-opt", align="center", do_wrap=False)

    def text(self, s, text, x, y, w, h, size=15, color=None, bold=False,
             mono=False, align="left", do_wrap=True, minimum=None):
        light = size >= 19 and not bold and not mono
        size = float(size)
        minimum = minimum if minimum is not None else size
        while True:
            lines = wrap(text, w, size, bold, mono, light) if do_wrap else text.split("\n")
            line_h = size / 72 * (1.26 if mono else 1.16)
            width_ok = max((measure(row, size, mono, bold, light) for row in lines), default=0) <= w + .005
            if (len(lines) * line_h <= h + .015 and width_ok) or size <= minimum:
                break
            size = max(minimum, size - .25)
        assert width_ok, f"Slide {self.i}: text too wide: {text!r}"
        assert len(lines) * line_h <= h + .02, f"Slide {self.i}: text too tall: {text!r}"
        sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        sh.name = text[:100].replace("\n", " / ")
        tf = sh.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for n, row in enumerate(lines):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = Pt(size * (1.26 if mono else 1.16))
            r = p.add_run()
            r.text = row
            r.font.name = "JetBrains Mono" if mono else "NVIDIA Sans Light" if light else "NVIDIA Sans"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color or C["ink"])
        self.audit.append(dict(slide=self.i, text=text, font_size=size, x=x, y=y, w=w, h=h))
        return sh

    def code(self, s, panel, x, y, w, h, size=11.5, green=False):
        code = panel["code"].strip("\n")
        rows = code.splitlines()
        # Keep code exactly on authored lines, without PowerPoint auto-wrapping.
        while (max(measure(row, size, mono=True, bold=True) for row in rows) > w - .32
               or len(rows) * size / 72 * 1.28 > h - .62) and size > 9.5:
            size -= .25
        assert max(measure(row, size, mono=True, bold=True) for row in rows) <= w - .30, (self.i, code)
        assert len(rows) * size / 72 * 1.28 <= h - .57, (self.i, code)
        self.rect(s, x, y, w, h, C["lightgreen"] if green else C["pale"])
        # Code provenance belongs in a distinct header, away from the source.
        self.rect(s, x, y, w, .40, C["green"])
        self.text(s, panel["label"], x + .17, y + .105, w - .34, .24,
                  10, C["ink"], True)
        cy = y + .53
        for row in rows:
            color = C["focus"] if any(v in row for v in ("@m.jit", "@m.kernel", "iter_args", "scf.yield")) else C["code"]
            bold = row.lstrip().startswith("@")
            self.text(s, row, x + .17, cy, w - .34, size / 72 * 1.30,
                      size, color, bold=bold, mono=True, do_wrap=False)
            cy += size / 72 * 1.28

    def base(self, data, start):
        self.i += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(C["white"])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        for logo in self.logos:
            el = deepcopy(logo)
            el.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", str(s.shapes._next_shape_id))
            s.shapes._spTree.insert_element_before(el, "p:extLst")
        special = data["layout"] in {"cover", "handoff", "placeholder"}
        if not special:
            self.text(s, data["title"], .48, .43, 9.04, .51, 21, minimum=19, do_wrap=False)
            if data.get("subtitle"):
                self.text(s, data["subtitle"], .50, 1.00, 9.02, .51, 11.8, C["focus"], minimum=11)
        number = data.get("display_number", self.i)
        if number:
            self.text(s, str(number).zfill(2), 8.8, 5.14, .7, .22, 8.5, C["gray"], align="right")
        if not special:
            self.text(s, data.get("section", ""), 5.5, 5.16, 3.1, .20, 8, C["gray"], align="right")
        if data.get("takeaway") and not special and data["layout"] not in {"principles", "compiler_pipelines", "design_simple"}:
            self.text(s, data["takeaway"], .50, 4.72, 9.00, .34, 11.5, C["gray"], minimum=10.5)
        label = str(number).zfill(2) if number else "TITLE"
        note = (f"SLIDE {label} — {data['title']}\nPDF page: {self.i}\n"
                f"Owner: {data['owner']}\n"
                f"Clock: {clock(start)}–{clock(start + data['seconds'])}\n"
                f"Target duration: {clock(data['seconds'])}\n\n"
                f"{data['notes']}\n\nSOURCES / PREPARATION\n" + "\n".join(data.get("sources", [])))
        s.notes_slide.notes_text_frame.text = note
        return s, note

    def render(self, d, start):
        s, note = self.base(d, start)
        layout = d["layout"]
        items = d.get("items", [])
        if render_custom(self, s, d):
            pass
        elif layout == "principles":
            for i, item in enumerate(items):
                y = 1.65 + i * .63
                self.text(s, str(i + 1) + ".", .55, y, .45, .42, 17, C["focus"], True)
                self.text(s, item["title"], 1.12, y, 8.33, .50, 17, minimum=16)
                if item.get("body"):
                    self.text(s, item["body"], 1.12, y + .39, 8.33, .32, 12, C["gray"])
            sh = self.text(s, d["takeaway"], .53, 4.43, 8.97, .45, 17, C["focus"])
            prefix, phrase = d["takeaway"].split(d["bold_phrase"])
            assert not phrase, "The emphasized takeaway phrase must end the sentence"
            p = sh.text_frame.paragraphs[0]
            p.clear()
            for value, bold in ((prefix, False), (d["bold_phrase"], True)):
                run = p.add_run()
                run.text = value
                run.font.name = "NVIDIA Sans"
                run.font.size = Pt(17)
                run.font.bold = bold
                run.font.color.rgb = RGBColor.from_string(C["focus"])
        elif layout == "compiler_pipelines":
            for r, pipeline in enumerate(d["pipelines"]):
                title_y = (1.47, 2.52, 3.94)[r]
                icon_y = (1.78, 2.86, 4.25)[r]
                self.text(s, pipeline["title"], .50, title_y, 4.5, .29, 12.5,
                          C["focus"] if r == 2 else C["ink"], True)
                n = len(pipeline["nodes"])
                if r < 2:
                    centers = [.98, 3.00, 5.02]
                else:
                    centers = [.98 + i * (8.04 / (n - 1)) for i in range(n)]
                for i, (node, cx) in enumerate(zip(pipeline["nodes"], centers)):
                    self.pipeline_node(s, node, cx, icon_y, label_width=1.60, accent=r == 2)
                    if i + 1 < n:
                        color = C["focus"] if r == 2 or i in pipeline.get("generation_edges", []) else C["gray"]
                        self.arrow(s, cx + .27, icon_y + .19, centers[i + 1] - .27, icon_y + .19, color)
                    if i in pipeline.get("generation_edges", []):
                        mid = (cx + centers[i + 1]) / 2
                        self.text(s, "run / generate", mid - .51, icon_y - .25, 1.02, .20,
                                  8.5, C["focus"], align="center")
                if pipeline.get("loop"):
                    # Route the repetition arrow around the nodes and their labels.
                    right = centers[-1] + .71
                    self.arrow(s, centers[-1] + .27, icon_y + .19, right, icon_y + .19, head=False)
                    self.arrow(s, right, icon_y + .19, right, 3.62, head=False)
                    self.arrow(s, right, 3.62, .42, 3.62, head=False)
                    self.arrow(s, .42, 3.62, .42, icon_y + .19, head=False)
                    self.arrow(s, .42, icon_y + .19, centers[0] - .27, icon_y + .19)
                    self.text(s, pipeline["loop_caption"], .70, 3.68, 4.80, .24, 10.5, C["focus"], align="center")
                    self.text(s, "…", 6.13, icon_y + .05, .52, .40, 21, C["gray"], align="center")
                    self.text(s, pipeline["repeat_caption"], 6.88, icon_y + .11, 2.62, .54, 12, C["gray"])
                if pipeline.get("final_edge_label"):
                    mid = (centers[-2] + centers[-1]) / 2
                    self.text(s, pipeline["final_edge_label"], mid - .86, icon_y - .18, 1.72, .21,
                              9, C["focus"], align="center")
        elif layout == "staging_benefits":
            # Open columns keep the two benefits readable without enclosing cards.
            self.rect(s, 4.92, 1.68, .012, 1.48, C["line"])
            for i, item in enumerate(items):
                x = .55 if i == 0 else 5.25
                self.text(s, f"0{i + 1}", x, 1.70, .34, .25, 10.5, C["focus"], True)
                self.text(s, item["label"], x + .48, 1.70, 3.72, .25, 10, C["focus"], True)
                self.text(s, item["title"], x, 2.12, 4.20, .47, 21)
                self.text(s, item["body"], x, 2.78, 4.20, .58, 12.5, C["gray"])
            self.rect(s, .55, 3.48, 8.90, .012, C["line"])
            self.text(s, d["example_context"], .55, 3.61, 8.90, .26, 10, C["gray"])
            for i, item in enumerate(items):
                x = .55 if i == 0 else 5.25
                color = C["gray"] if i == 0 else C["focus"]
                self.text(s, item["expression"], x, 4.02, 1.62, .35, 14, color, mono=True)
                self.arrow(s, x + 1.77, 4.18, x + 2.20, 4.18, color)
                self.text(s, item["result"], x + 2.41, 4.01, 1.85, .33, 14, color,
                          mono=i == 1, bold=i == 0)
                self.text(s, item["caption"], x + 2.41, 4.38, 1.85, .25, 10, C["gray"])
        elif layout == "stages":
            for i, item in enumerate(items):
                x = .52 if i == 0 else 5.32
                self.text(s, item["title"], x, 1.85, 4.16, .60, 18, C["ink"], minimum=16)
                self.text(s, item["label"], x, 2.55, 4.16, .40, 13, C["focus"], True)
                self.text(s, item["body"], x, 3.11, 4.16, 1.18, 13, C["gray"], minimum=12)
            self.arrow(s, 4.71, 2.68, 5.11, 2.68)
            self.text(s, "compile", 4.61, 2.91, .65, .22, 9.2, C["gray"], align="center")
        elif layout == "agenda":
            for i, item in enumerate(items):
                y = 1.70 + i * .68
                self.text(s, f"{i + 1:02}", .55, y + .025, .45, .35, 13, C["focus"], True)
                self.text(s, item["title"], 1.16, y, 8.30, .44, 17, minimum=16)
        elif layout == "cover":
            self.text(s, d.get("event", ""), .52, .58, 8.96, .30, 12, C["gray"])
            self.text(s, d.get("date", ""), .52, .99, 8.96, .28, 11.5, C["gray"])
            self.text(s, d["display_title"], .48, 1.81, 9.04, 1.23, 31, minimum=29)
            if d.get("subtitle"):
                self.text(s, d["subtitle"], .52, 3.27, 8.98, .40, 14, C["focus"], minimum=13)
            for index, presenter in enumerate(d.get("presenters", [])):
                x = .52 + index * 4.62
                self.text(s, presenter["name"], x, 4.11, 4.20, .36, 18)
                if presenter.get("role"):
                    self.text(s, presenter["role"], x, 4.58, 4.20, .44,
                              11.8, C["gray"], minimum=11)
        elif layout in {"handoff", "placeholder"}:
            self.text(s, d.get("kicker", "NEXT"), .52, .72, 8.95, .35, 12, C["focus"], True)
            self.text(s, d["title"], .49, 1.78, 9, 1.35, 28, minimum=23)
            self.text(s, d.get("subtitle", ""), .52, 3.41, 8.96, .88, 17, C["gray"], minimum=15)
            self.rect(s, .52, 4.56, .48, .06, C["green"])
            self.text(s, d.get("footer", d.get("owner", "")), 1.15, 4.47, 8, .40, 12, C["focus"])
        elif layout in {"cards", "compare"}:
            n = len(items)
            gap = .22
            w = (9.04 - (n - 1) * gap) / n
            for i, item in enumerate(items):
                x = .48 + i * (w + gap)
                highlight = layout == "compare" and i == n - 1
                self.rect(s, x, 1.72, w, 2.69, C["lightgreen"] if highlight else C["pale"])
                self.rect(s, x, 1.72, w, .045, C["green"] if highlight else C["line"])
                self.text(s, item["title"], x + .20, 1.94, w - .40, .63, 17,
                          C["focus"] if highlight else C["ink"], True, minimum=14)
                self.text(s, item["body"], x + .20, 2.82, w - .40, 1.32, 13, C["gray"], minimum=11)
        elif layout == "code_pair":
            self.code(s, d["left"], .48, 1.65, 4.40, 2.84, size=11.3)
            self.code(s, d["right"], 5.10, 1.65, 4.42, 2.84, size=11.3, green=True)
        elif layout == "code_notes":
            self.code(s, d["left"], .48, 1.65, 5.70, 2.84, size=11.6)
            n = len(items)
            height = 2.84 / n
            for i, item in enumerate(items):
                y = 1.69 + i * height
                self.text(s, item["title"], 6.47, y, 3.03, .44, 14, C["focus"], True, minimum=12)
                self.text(s, item["body"], 6.47, y + .48, 3.03, height - .53, 12, C["gray"], minimum=10.5)
        elif layout == "flow":
            n = len(items)
            gap = .25
            w = (9.04 - (n - 1) * gap) / n
            for i, item in enumerate(items):
                x = .48 + i * (w + gap)
                self.rect(s, x, 2.03, w, 1.93, C["lightgreen"] if i in {0, n - 1} else C["pale"], C["line"])
                self.text(s, item["title"], x + .14, 2.23, w - .28, .62, 15, C["focus"], True, minimum=12)
                self.text(s, item["body"], x + .14, 3.03, w - .28, .73, 11.5, C["gray"], minimum=10)
                if i < n - 1:
                    self.arrow(s, x + w + .035, 2.96, x + w + gap - .035, 2.96)
            if d.get("footnote"):
                self.text(s, d["footnote"], .50, 4.18, 9.0, .35, 11.5, C["gray"], minimum=10.5)
        elif layout == "table":
            headers, rows = d["headers"], d["rows"]
            n = len(headers)
            widths = [9.04 / n] * n
            if n == 3:
                widths = [2.05, 3.45, 3.54]
            elif n == 2:
                widths = [3.04, 6.0]
            self.rect(s, .48, 1.66, 9.04, .48, C["emphasis"])
            x = .48
            for head, w in zip(headers, widths):
                self.text(s, head, x + .15, 1.77, w - .3, .29, 11.5, C["focus"], True)
                x += w
            rh = min(.83, 2.34 / len(rows))
            for r, row in enumerate(rows):
                y = 2.14 + r * rh
                self.rect(s, .48, y, 9.04, rh, C["pale"] if r % 2 == 0 else C["white"])
                x = .48
                for col, (cell, w) in enumerate(zip(row, widths)):
                    self.text(s, cell, x + .15, y + .13, w - .30, rh - .20, 12,
                              C["ink"] if col == 0 else C["gray"], bold=col == 0, minimum=10)
                    x += w
        else:
            raise ValueError(layout)
        return note


def write_outline(data, notes):
    start = 0
    entries = []
    with (HERE / "slide-skeleton.csv").open("w", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(["slide", "pdf_page", "start", "end", "duration", "speaker", "section", "title", "teaching_point"])
        for i, d in enumerate(data["slides"], 1):
            end = start + d["seconds"]
            label = d.get("display_number", i) or "Title"
            writer.writerow([label, i, clock(start), clock(end), clock(d["seconds"]), d["owner"], d["section"], d["title"], d.get("takeaway", "")])
            entries.append(f'<tr><td>{label}<br><span>PDF {i}</span></td><td>{clock(start)}–{clock(end)}</td><td>{html.escape(d["owner"])}</td><td><b>{html.escape(d["title"])}</b><br><span>{html.escape(d.get("takeaway", ""))}</span><details><summary>Speaker notes and sources</summary><pre>{html.escape(notes[i-1])}</pre></details></td></tr>')
            start = end
    total_seconds = sum(slide["seconds"] for slide in data["slides"])
    phase_summary = "; ".join(f"{phase}: {clock(seconds)}"
                              for phase, seconds in data["timing"].items())
    title = html.escape(data["title"])
    page = '''<!doctype html><html lang="en"><meta charset="utf-8"><title>mlir dsl tutorial · slide skeleton</title>
<style>body{font:16px/1.5 system-ui,sans-serif;max-width:1120px;margin:48px auto;padding:0 24px;color:#202020}h1{font-weight:400;font-size:36px}a{color:#416600}table{border-collapse:collapse;width:100%;font-size:14px}th{background:#e4f0d0;text-align:left}th,td{padding:14px;border-bottom:1px solid #ddd;vertical-align:top}td:nth-child(2){white-space:nowrap}span{color:#616161}summary{cursor:pointer;color:#416600;margin-top:8px}pre{white-space:pre-wrap;font:14px/1.5 system-ui}p{max-width:950px}</style>
'''
    page += f'<h1>{title}</h1>\n'
    page += (f'<p>Draft v16 · Guray Ozen and Amir Tavakkoli · '
             f'LLVM/MLIR developers new to Python DSLs · {total_seconds / 60:g} minutes total. '
             f'Preparation timing: {html.escape(phase_summary)}. '
             'The mutation placeholder is reserved for Amir’s slides; Guray returns for the sub-DSL showcase.</p>\n')
    page += (f'<p><a href="{NAME}.pptx">Editable PowerPoint</a> · '
             f'<a href="{NAME}.pdf">PDF preview</a> · '
             '<a href="overview.png">All slides</a> · '
             '<a href="slide-skeleton.csv">Timing CSV</a> · '
             '<a href="verification/v15/scalar-bindings.txt">Scalar bindings check</a></p>\n')
    page += ('<p>Examples follow mlir dsl in this checkout. Speaker notes identify '
             'source-checked code, executed examples, and proposed target adapters. '
             'PyIR content is reserved for Amir.</p>\n'
             '<table><thead><tr><th>#</th><th>Clock</th><th>Speaker</th>'
             '<th>Slide / teaching point</th></tr></thead><tbody>')
    page += "".join(entries) + "</tbody></table></html>"
    (HERE / "slide-skeleton.html").write_text(page)
    (HERE / "speaker-notes.txt").write_text("\n\n" + ("\n\n" + "=" * 76 + "\n\n").join(notes) + "\n")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args()
    data = json.loads((HERE / "deck.json").read_text())
    assert set(d["phase"] for d in data["slides"]) <= set(data["timing"])
    for phase, duration in data["timing"].items():
        actual = sum(d["seconds"] for d in data["slides"] if d["phase"] == phase)
        assert actual == duration, (phase, actual, duration)
    deck = Deck(args.template)
    notes = []
    start = 0
    for d in data["slides"]:
        notes.append(deck.render(d, start))
        start += d["seconds"]
    deck.prs.core_properties.title = data["title"]
    phase_summary = "; ".join(f"{phase}: {seconds / 60:g} minutes"
                              for phase, seconds in data["timing"].items())
    deck.prs.core_properties.subject = f"LLVM conference tutorial · mlir dsl · {phase_summary}"
    deck.prs.core_properties.author = "Guray Ozen; Amir Tavakkoli"
    deck.prs.core_properties.keywords = "MLIR, Python, DSL, tutorial, draft"
    deck.prs.core_properties.comments = "Draft v16. Uses the supplied cutlass-python-pytorch-v22 template and style."
    deck.prs.save(HERE / f"{NAME}.pptx")
    for shape in deck.audit:
        assert shape["x"] >= 0 and shape["y"] >= 0
        assert shape["x"] + shape["w"] <= 10.001
        assert shape["y"] + shape["h"] <= 5.626
    (HERE / "layout-audit.json").write_text(json.dumps(deck.audit, indent=2))
    write_outline(data, notes)
    print(f"Built {len(data['slides'])} slides: {phase_summary}")
    print(HERE / f"{NAME}.pptx")


if __name__ == "__main__":
    main()
