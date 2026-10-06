"""Native PowerPoint diagrams for the mlir dsl tutorial.

The caller supplies the Deck drawing primitives, keeping the template, fonts,
color palette, and measured text audit shared with the rest of the deck.
"""

from pptx.dml.color import RGBColor
from pptx.util import Pt


C = dict(ink="000000", gray="616161", line="D5D5D5", pale="F7F7F7",
         white="FFFFFF", green="76B900", code="484848", focus="416600",
         emphasis="E4F0D0", lightgreen="F0F6E6")


def _statement(deck, slide, value, x, y, w, h, size=14,
               bold_phrase="multi-stage programming"):
    """Emphasize a phrase without rasterizing or losing editable text."""
    shape = deck.text(slide, value, x, y, w, h, size, C["focus"])
    for paragraph in shape.text_frame.paragraphs:
        line = paragraph.text
        if bold_phrase not in line:
            continue
        before, after = line.split(bold_phrase, 1)
        paragraph.clear()
        for text, bold in ((before, False), (bold_phrase, True), (after, False)):
            run = paragraph.add_run()
            run.text = text
            run.font.name = "NVIDIA Sans"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(C["focus"])
    return shape


def _box(deck, slide, text, x, y, w, h, *, accent=False, size=12):
    deck.rect(slide, x, y, w, h,
              C["lightgreen"] if accent else C["pale"], C["line"])
    deck.text(slide, text, x + .12, y + .08, w - .24, h - .16,
              size, C["focus"] if accent else C["ink"],
              bold=accent, align="center", minimum=size - 1)


def _flow(deck, slide, labels, x, y, w, size=10.8):
    if isinstance(labels, str):
        labels = [part.strip() for part in labels.replace("->", "→").split("→")]
    if not labels:
        return
    gap = .30
    node_w = (w - gap * (len(labels) - 1)) / len(labels)
    for index, label in enumerate(labels):
        nx = x + index * (node_w + gap)
        deck.text(slide, label, nx, y, node_w, .45, size,
                  C["focus"] if index == len(labels) - 1 else C["gray"],
                  align="center", minimum=9.5)
        if index < len(labels) - 1:
            deck.arrow(slide, nx + node_w + .035, y + .15,
                       nx + node_w + gap - .035, y + .15)


def render_custom(deck, slide, data):
    """Render a custom layout, returning False for a layout owned by Deck."""
    layout = data["layout"]
    items = data.get("items", [])

    if layout == "use_cases":
        # Each use case is a small explanatory picture. The visual shows where
        # Python enters the system, rather than pairing an icon with prose.
        for index, item in enumerate(items[:4]):
            x = .53 + (index % 2) * 4.61
            y = 1.43 + (index // 2) * 1.65
            width, height = 4.33, 1.55
            deck.rect(slide, x, y, width, height, C["pale"])
            deck.rect(slide, x, y, .045, height, C["green"])
            # Numbering provides an immediate reading order for the four
            # pictures, with a compact badge on the shared title baseline.
            deck.rect(slide, x + .17, y + .105, .35, .35, C["green"])
            deck.text(slide, f"{index + 1:02}", x + .18, y + .185,
                      .33, .22, 11.5, C["ink"], True, align="center")
            deck.text(slide, item["title"], x + .66, y + .12,
                      width - .84, .33, 16.5, minimum=15.5)
            if index == 0:
                labels = item.get("diagram_labels", ["Public", "Internal", "LLVM"])
                node_x = (x + .22, x + 1.61, x + 3.11)
                node_w = (.95, 1.13, .98)
                for n, (nx, nw, label) in enumerate(zip(node_x, node_w, labels)):
                    deck.rect(slide, nx, y + 1.00, nw, .36,
                              C["lightgreen"] if n == 1 else C["white"],
                              C["green"] if n == 1 else C["line"])
                    deck.text(slide, label, nx + .025, y + 1.075,
                              nw - .05, .24, 12.2,
                              C["focus"] if n == 1 else C["gray"],
                              bold=n == 1, align="center", minimum=11.5)
                    if n < 2:
                        deck.arrow(slide, nx + nw + .035, y + 1.18,
                                   node_x[n + 1] - .045, y + 1.18)
                # Testing enters directly at the normally internal dialect.
                deck.rect(slide, x + 1.57, y + .55, 1.21, .28, C["green"])
                deck.text(slide, "Python test", x + 1.61, y + .605,
                          1.13, .19, 10.5, C["ink"], True, align="center")
                deck.arrow(slide, x + 2.175, y + .845,
                           x + 2.175, y + .96)
            elif index in (1, 3):
                labels = item.get("diagram_labels", ["Python", "Recipe IR", "External\nconsumer"]
                                  if index == 1 else ["Python code", "DSL compiler", "GPU"])
                kinds = ("source", "source", "compiler") if index == 1 else ("source", "compiler", "binary")
                centers = [x + .73, x + 2.14, x + 3.55]
                for n, (center, label, kind) in enumerate(zip(centers, labels, kinds)):
                    deck.pipeline_node(slide, dict(kind=kind, label=label),
                                       center, y + .51, label_width=1.31,
                                       accent=n == 1 or index == 3 and n == 2)
                    if n < 2:
                        deck.arrow(slide, center + .28, y + .70,
                                   centers[n + 1] - .29, y + .70)
                if index == 3 and item.get("link"):
                    link = item["link"]
                    shape = deck.text(slide, "NVIDIA CuTe DSL / CUTLASS Python",
                                      x + .18, y + 1.35, width - .36, .18,
                                      9.4, C["focus"], align="center")
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.hyperlink.address = link["url"]
                            run.font.underline = True
            else:
                # Existing modules plug directly into one downstream DSL.
                labels = item.get("diagram_labels", ["JIT", "Cache", "Env"])
                deck.rect(slide, x + 1.98, y + .53, 2.11, .91,
                          C["white"], C["green"])
                deck.text(slide, "MyDSL", x + 2.60, y + .86,
                          1.23, .31, 17, C["focus"], True, align="center")
                for n, label in enumerate(labels[:3]):
                    by = y + .55 + n * .31
                    deck.rect(slide, x + .25, by, 1.12, .25,
                              C["lightgreen"], C["green"])
                    deck.text(slide, label, x + .285, by + .043,
                              1.05, .19, 10.7, C["focus"], True, align="center")
                    deck.arrow(slide, x + 1.42, by + .125,
                               x + 1.95, by + .125)
                    deck.rect(slide, x + 2.11, by + .035, .20, .18,
                              C["green"])

    elif layout == "architecture":
        # The top boxes are thin sub-DSLs. All arrows share one reusable core.
        subdsls = data.get("subdsls", [
            dict(title="CPU DSL", subtitle="arith / llvm"),
            dict(title="GPU DSL", subtitle="gpu / nvvm"),
            dict(title="Your DSL", subtitle="your dialect"),
        ])
        node_w, gap, x0 = 2.66, .48, .53
        centers = [x0 + index * (node_w + gap) + node_w / 2
                   for index in range(len(subdsls[:3]))]
        rail_y, base_y = 3.02, 3.40
        deck.arrow(slide, 5.0, base_y, 5.0, rail_y, head=False)
        if centers:
            deck.arrow(slide, centers[0], rail_y, centers[-1], rail_y,
                       head=False)
        for index, item in enumerate(subdsls[:3]):
            x = x0 + index * (node_w + gap)
            center = x + node_w / 2
            deck.arrow(slide, center, rail_y, center, 2.65)
            deck.rect(slide, x, 1.67, node_w, .94,
                      C["lightgreen"], C["green"])
            deck.text(slide, item["title"], x + .12, 1.86,
                      node_w - .24, .35, 20, C["focus"], align="center")
            deck.text(slide, item.get("subtitle", item.get("body", "")), x + .12, 2.28,
                      node_w - .24, .25, 12.5, C["gray"], align="center")
        deck.rect(slide, .53, base_y, 8.94, 1.16,
                  C["pale"], C["line"])
        deck.text(slide, data.get("core_label", "BaseDSL"), .75, 3.53,
                  8.50, .33, 18, C["focus"], True, align="center")
        labels = data.get("core_labels", [])
        gap, left, total = .10, .74, 8.52
        chip_w = (total - gap * (len(labels) - 1)) / max(len(labels), 1)
        for index, label in enumerate(labels):
            x = left + index * (chip_w + gap)
            deck.rect(slide, x, 4.02, chip_w, .36, C["white"])
            deck.text(slide, label, x + .025, 4.085,
                      chip_w - .05, .23, 10.8, C["ink"],
                      align="center", minimum=10.3)

    elif layout == "opening_text":
        # These opening beats need typography, not another diagram. The lack
        # of cards or connectors gives the audience a moment to read the idea.
        lead = data.get("lead", "")
        if lead:
            deck.text(slide, lead, .59, 1.65, 8.80, .80,
                      23, C["ink"], minimum=21)
        start, step = (2.81, .65) if lead else (1.72, .74)
        for index, line in enumerate(data.get("lines", [])[:3]):
            deck.text(slide, line, .61, start + index * step,
                      8.73, .50, 20 if lead else 27,
                      C["gray"] if lead else C["ink"],
                      minimum=19 if lead else 25)
        if data.get("footer"):
            deck.text(slide, data["footer"], .61, 4.34, 8.73, .44,
                      16, C["focus"], minimum=15)

    elif layout == "plugin_assembly":
        # Selected modules attach directly to a portable core. One composed
        # system produces one language namespace, with no hierarchy fan-out.
        plugins = data.get("plugins", [
            dict(title="Dialect", body="Types + passes"),
            dict(title="AST capture", body="Control flow"),
            dict(title="Interop", body="Tensors"),
        ])
        deck.text(slide, data.get("caption", "Choose the plugins your DSL needs"),
                  .58, 1.33, 6.14, .31, 14, C["focus"], minimum=13)
        count, gap, total = len(plugins[:3]), .24, 6.08
        module_w = (total - gap * (count - 1)) / max(count, 1)
        for index, plugin in enumerate(plugins[:3]):
            x = .53 + index * (module_w + gap)
            center = x + module_w / 2
            # A physical tab connects each chosen plugin to the shared core.
            deck.rect(slide, center - .095, 2.59, .19, .29, C["green"])
            deck.rect(slide, x, 1.85, module_w, .83,
                      C["lightgreen"], C["green"])
            deck.text(slide, plugin["title"], x + .10, 1.98,
                      module_w - .20, .30, 15.5, C["focus"], True,
                      align="center", minimum=14.5)
            if plugin.get("body"):
                deck.text(slide, plugin["body"], x + .10, 2.39,
                          module_w - .20, .24, 11, C["gray"],
                          align="center", minimum=10.5)
        deck.rect(slide, .53, 2.86, 6.08, 1.42, C["pale"], C["line"])
        deck.text(slide, data.get("core_label", "mlir dsl core"),
                  .77, 3.10, 5.60, .43, 21, C["ink"], True,
                  align="center", minimum=19.5)
        features = data.get("core_features", ["Staging", "Types", "JIT + cache", "Diagnostics"])
        chip_gap, chip_total = .10, 5.64
        chip_w = (chip_total - chip_gap * (len(features) - 1)) / max(len(features), 1)
        for index, feature in enumerate(features):
            x = .75 + index * (chip_w + chip_gap)
            deck.rect(slide, x, 3.79, chip_w, .32, C["white"])
            deck.text(slide, feature, x + .03, 3.855,
                      chip_w - .06, .22, 11.5, C["gray"],
                      align="center", minimum=10.8)
        deck.arrow(slide, 6.72, 3.53, 7.22, 3.53)
        result = data.get("result", dict(title="Your DSL", body="your namespace"))
        deck.rect(slide, 7.34, 3.06, 2.13, .96,
                  C["lightgreen"], C["green"])
        deck.text(slide, result["title"], 7.48, 3.22,
                  1.85, .33, 19, C["focus"], align="center", minimum=17)
        deck.text(slide, result.get("body", ""), 7.46, 3.70,
                  1.89, .23, 11, C["gray"], align="center", minimum=10.5)

    elif layout == "core_feature_boxes":
        # Infrastructure is the subject of this slide. The enclosing box and
        # four substantial feature boxes make the reusable boundary explicit.
        deck.rect(slide, .53, 1.48, 8.94, 3.10, C["pale"], C["line"])
        deck.text(slide, data.get("core_label", "Core DSL"),
                  .79, 1.70, 8.42, .39, 22, C["ink"], True,
                  minimum=21)
        features = data.get("core_features", [
            dict(title="Diagnostics", detail="Errors · remarks"),
            dict(title="AST preprocessing", detail="Plugin hooks"),
            dict(title="JIT cache", detail="Memory · disk"),
            dict(title="Type inference", detail="Promotion · casts"),
        ])
        for index, feature in enumerate(features[:4]):
            x = .79 + index % 2 * 4.29
            y = 2.25 + index // 2 * 1.08
            width, height = 4.10, .95
            deck.rect(slide, x, y, width, height, C["white"], C["line"])
            deck.rect(slide, x, y, .055, height, C["green"])
            deck.text(slide, feature["title"], x + .20, y + .16,
                      width - .40, .36, 20, C["focus"], True,
                      minimum=18.5, do_wrap=False)
            deck.text(slide, feature.get("detail", ""), x + .20, y + .65,
                      width - .40, .23, 13.1, C["gray"], minimum=12.4)
        if data.get("core_note"):
            deck.text(slide, data["core_note"], .59, 4.85, 8.82, .28,
                      13.6, C["gray"], align="center", minimum=12.8)

    elif layout == "plugin_boxes":
        # A shelf of plugin families can be reused in either composition.
        # Both results visibly contain the same neutral core foundation.
        plugins = data.get("plugins", [
            dict(title="Dialect", examples=["LLVM", "SCF", "GPU"]),
            dict(title="AST preprocessor", examples=["SCF capture"]),
            dict(title="Third party", examples=["DLPack", "TVM-FFI"]),
        ])
        plugin_centers = []
        for index, plugin in enumerate(plugins[:3]):
            x, y, width, height = .53 + index * 3.05, 1.48, 2.84, 1.12
            center = x + width / 2
            plugin_centers.append(center)
            deck.rect(slide, x, y, width, height, C["lightgreen"], C["green"])
            deck.rect(slide, x + .15, y + .18, .32, .32, C["green"])
            deck.text(slide, str(index + 1), x + .17, y + .244,
                      .28, .22, 11.5, C["ink"], True, align="center")
            deck.text(slide, plugin["title"], x + .60, y + .19,
                      width - .75, .32, 15.4, C["focus"], True,
                      minimum=14, do_wrap=False)
            examples = plugin.get("examples", [])
            if isinstance(examples, str):
                examples = [s.strip() for s in examples.split("·")]
            gap, chip_total = .09, width - .30
            chip_w = (chip_total - gap * (len(examples) - 1)) / max(len(examples), 1)
            for number, example in enumerate(examples):
                bx = x + .15 + number * (chip_w + gap)
                deck.rect(slide, bx, y + .70, chip_w, .28, C["white"])
                deck.text(slide, example, bx + .03, y + .753,
                          chip_w - .06, .21, 10.8, C["focus"],
                          align="center", minimum=10)
            deck.arrow(slide, center, y + height, center, 2.85, head=False)
        if plugin_centers:
            deck.arrow(slide, plugin_centers[0], 2.85,
                       plugin_centers[-1], 2.85, head=False)
        dsls = data.get("dsls", [
            dict(title="mlir dsl", subtitle="Working end-to-end",
                 plugins=["SCF", "LLVM", "SCF AST"]),
            dict(title="Your DSL", subtitle="Reuse + extend",
                 plugins=["Existing plugins", "Your additions"]),
        ])
        for index, dsl in enumerate(dsls[:2]):
            x, y, width, height = .53 + index * 4.61, 3.19, 4.33, 1.59
            center = x + width / 2
            deck.arrow(slide, center, 2.85, center, y - .035)
            deck.rect(slide, x, y, width, height, C["white"], C["green"])
            deck.text(slide, dsl["title"], x + .17, y + .14,
                      width - .34, .36, 20, C["focus"], True,
                      align="center", minimum=19)
            deck.text(slide, dsl.get("subtitle", ""), x + .17, y + .60,
                      width - .34, .23, 11.5, C["gray"],
                      align="center", minimum=10.8)
            selected = dsl.get("plugins", ["Selected plugins"])
            chip_gap, chip_total = .10, width - .34
            chip_w = (chip_total - chip_gap * (len(selected) - 1)) / max(len(selected), 1)
            for number, label in enumerate(selected):
                bx = x + .17 + number * (chip_w + chip_gap)
                deck.rect(slide, bx, y + .97, chip_w, .27, C["lightgreen"])
                deck.text(slide, label, bx + .035, y + 1.018,
                          chip_w - .07, .22, 11.1, C["focus"],
                          align="center", minimum=10.5)
            deck.rect(slide, x + .17, y + 1.25, width - .34, .24, C["pale"])
            deck.text(slide, data.get("core_label", "Core DSL"),
                      x + .20, y + 1.286, width - .40, .20,
                      10.8, C["ink"], True, align="center")
        if data.get("caption"):
            deck.text(slide, data["caption"], .59, 4.93, 8.82, .27,
                      13, C["focus"], align="center", minimum=12.4)

    elif layout == "plugin_behavior":
        # Explain extension points by showing the work done at each hook.
        # The core is a slim service strip; the rows show concrete behavior,
        # rather than repeating the architecture as three oversized boxes.
        deck.text(slide, data.get("core_label", "Shared core"),
                  .59, 1.36, 4.21, .30, 14.4, C["ink"], True)
        if data.get("core_note"):
            deck.text(slide, data["core_note"], 4.89, 1.385,
                      4.46, .25, 11.5, C["gray"], align="right")
        deck.rect(slide, .53, 1.78, 8.94, .33, C["pale"])
        features = data.get("core_features", [
            "Diagnostics", "AST hooks", "JIT cache", "Type inference",
        ])
        for index, feature in enumerate(features[:4]):
            x = .70 + index * 2.23
            deck.text(slide, feature, x, 1.827, 2.05, .25,
                      12.6, C["gray"], align="center", minimum=11.8)
        plugins = data.get("plugins", [])
        for index, plugin in enumerate(plugins[:3]):
            y = 2.21 + index * .84
            deck.rect(slide, .57, y + .08, .33, .33, C["green"])
            deck.text(slide, str(index + 1), .59, y + .145,
                      .29, .21, 11.5, C["ink"], True, align="center")
            deck.text(slide, plugin["title"], 1.04, y + .045,
                      2.07, .30, 15.3, C["focus"], True, minimum=14.3)
            deck.text(slide, plugin.get("role", ""), 1.04, y + .445,
                      2.07, .27, 11.8, C["gray"], minimum=10.8)
            # Input and output remain plain editable code. Only the actual
            # hook has a light highlight, so the action is the focal point.
            input_value = plugin.get("input", "")
            input_y = y + (.10 if "\n" in input_value else .25)
            deck.text(slide, input_value, 3.24, input_y, 1.70, .50,
                      12.4, C["code"], mono=True, align="center",
                      minimum=9.8, do_wrap=False)
            deck.arrow(slide, 4.98, y + .35, 5.25, y + .35)
            deck.rect(slide, 5.32, y + .085, 1.77, .54, C["lightgreen"])
            deck.text(slide, plugin.get("hook", ""), 5.40, y + .215,
                      1.61, .33, 12.4, C["focus"], True,
                      align="center", minimum=10.6, do_wrap=False)
            deck.arrow(slide, 7.15, y + .35, 7.42, y + .35)
            deck.text(slide, plugin.get("output", ""), 7.51, y + .215,
                      1.80, .35, 14, C["focus"], mono=True,
                      align="center", minimum=12.5, do_wrap=False)
            if plugin.get("note"):
                deck.text(slide, plugin["note"], 3.25, y + .655,
                          6.03, .19, 9.2, C["gray"], align="center")
            if index < 2:
                deck.rect(slide, 1.04, y + .835, 8.27, .006, C["line"])
        footer = data.get("footer", data.get("core_note", ""))
        if footer:
            deck.text(slide, footer, .59, 4.89, 8.76, .29,
                      12.5, C["gray"], minimum=11.8)

    elif layout == "plugin_execution":
        # A runnable example makes composition tangible: the selected plugin
        # emits the IR, provides lowering, and produces the displayed result.
        deck.text(slide, data.get("recipe_label", "Reference composition"),
                  .57, 1.36, 2.36, .29, 13, C["gray"], minimum=12.3)
        deck.text(slide, data.get("recipe", "Core + LLVM + SCF + SCF AST"),
                  3.00, 1.32, 6.39, .37, 18, C["focus"], True, minimum=16.5)
        deck.text(slide, data.get("emission_label", "LLVM plugin: emit IR"),
                  4.36, 1.73, 3.34, .29, 12.5, C["focus"], minimum=11.8)
        lowering = data.get("lowering_label", "LLVM plugin: lower + run").replace(": ", "\n", 1)
        deck.text(slide, lowering,
                  8.02, 1.71, 1.44, .56, 11.6, C["focus"],
                  align="center", minimum=10.8)
        deck.code(slide, data["python"], .53, 2.08, 3.41, 1.94, size=10.9)
        deck.arrow(slide, 4.00, 3.12, 4.29, 3.12)
        deck.code(slide, data["ir"], 4.36, 2.08, 3.34, 1.94, size=10.2)
        deck.arrow(slide, 7.77, 3.12, 8.04, 3.12)
        deck.text(slide, data.get("result_label", "CPU result"),
                  8.10, 2.73, 1.29, .29, 12, C["gray"], align="center")
        deck.text(slide, str(data.get("result", "12")),
                  8.10, 3.20, 1.29, .76, 44, C["focus"], True,
                  align="center", minimum=41)
        deck.code(slide, data["reuse"], .53, 4.12, 8.94, .96, size=11.8)

    elif layout == "core_plugin_types":
        # Separate the infrastructure from the three extension families. The
        # short connections indicate hooks, not an executable lowering chain.
        deck.rect(slide, .53, 1.48, 8.94, 1.12, C["pale"], C["line"])
        deck.text(slide, data.get("core_label", "Shared core"),
                  .76, 1.66, 3.10, .36, 20, C["ink"], True)
        deck.text(slide, data.get("core_note", "The core is infrastructure."),
                  4.06, 1.72, 5.12, .30, 14, C["gray"],
                  align="right", minimum=13)
        features = data.get("core_features", [
            "Diagnostics", "AST-preprocess hooks", "JIT cache", "Type inference",
        ])
        feature_gap, feature_total = .15, 8.48
        feature_w = (feature_total - feature_gap * (len(features) - 1)) / max(len(features), 1)
        for index, feature in enumerate(features):
            x = .76 + index * (feature_w + feature_gap)
            deck.rect(slide, x, 2.17, feature_w, .29, C["white"])
            deck.text(slide, feature, x + .025, 2.217,
                      feature_w - .05, .24, 12.2, C["gray"],
                      align="center", minimum=11.5)
        plugins = data.get("plugins", [
            dict(title="Dialect", purpose="Types + operations", examples="LLVM · SCF · GPU"),
            dict(title="AST preprocessor", purpose="Python control flow", examples="SCF AST"),
            dict(title="Third party", purpose="Connect other tools", examples="PyTorch · DLPack · TVM-FFI"),
        ])
        for index, plugin in enumerate(plugins[:3]):
            x, width, y = .53 + index * 3.05, 2.84, 3.08
            center = x + width / 2
            deck.arrow(slide, center, y - .04, center, 2.65)
            deck.rect(slide, x, y, width, 1.57, C["lightgreen"], C["green"])
            deck.rect(slide, x + .16, y + .19, .32, .32, C["green"])
            deck.text(slide, str(index + 1), x + .18, y + .259,
                      .28, .20, 11.3, C["ink"], True, align="center")
            deck.text(slide, plugin["title"], x + .61, y + .22,
                      width - .76, .31, 15, C["focus"], True,
                      minimum=13.5, do_wrap=False)
            deck.text(slide, plugin.get("purpose", ""), x + .16, y + .75,
                      width - .32, .31, 14.3, C["ink"], minimum=13.2)
            deck.text(slide, plugin.get("examples", ""), x + .16, y + 1.22,
                      width - .32, .23, 11.1, C["focus"], minimum=10.4)

    elif layout == "dsl_composition":
        # A shared shelf can serve more than one complete DSL. Each result has
        # its own core + chosen plugin composition and a full execution path.
        deck.text(slide, data.get("library_label", "Reuse existing plugins"),
                  .60, 1.43, 8.80, .35, 17, C["focus"], True)
        deck.rect(slide, .53, 1.93, 8.94, .74, C["pale"], C["line"])
        groups = data.get("library_groups", [
            dict(title="Dialect", examples="LLVM · SCF · GPU"),
            dict(title="AST preprocessor", examples="SCF AST"),
            dict(title="Third party", examples="PyTorch · DLPack · TVM-FFI"),
        ])
        for index, group in enumerate(groups[:3]):
            x, width = .68 + index * 3.00, 2.66
            deck.text(slide, group["title"], x, 2.04, width, .27,
                      13.5, C["ink"], True, align="center")
            deck.text(slide, group.get("examples", ""), x, 2.42,
                      width, .23, 10.8, C["gray"], align="center", minimum=10.3)
        # Branching illustrates independent selection, not a requirement to
        # install every plugin on the shelf or invent a new plugin per DSL.
        deck.arrow(slide, 5.00, 2.68, 5.00, 2.86, head=False)
        deck.arrow(slide, 2.695, 2.86, 7.305, 2.86, head=False)
        dsls = data.get("dsls", [
            dict(title="mlir dsl", subtitle="Working end-to-end showcase",
                 composition="Core + selected plugins", pipeline=["Python", "MLIR", "Execute"]),
            dict(title="Your DSL", subtitle="Your end-to-end language",
                 composition="Core + existing + your plugins", pipeline=["Python", "Your IR", "Runtime"]),
        ])
        for index, dsl in enumerate(dsls[:2]):
            x, y, width = .53 + index * 4.61, 3.14, 4.33
            center = x + width / 2
            deck.arrow(slide, center, 2.86, center, y - .04)
            deck.rect(slide, x, y, width, 1.57, C["lightgreen"], C["green"])
            deck.text(slide, dsl["title"], x + .19, y + .16,
                      width - .38, .34, 19, C["focus"], True,
                      align="center", minimum=18)
            deck.text(slide, dsl.get("subtitle", ""), x + .18, y + .60,
                      width - .36, .24, 11.5, C["gray"], align="center")
            deck.rect(slide, x + .18, y + .93, width - .36, .28, C["white"])
            deck.text(slide, dsl.get("composition", "Core + selected plugins"),
                      x + .25, y + .981, width - .50, .23,
                      11.5, C["ink"], align="center", minimum=11)
            _flow(deck, slide, dsl.get("pipeline", ["Python", "MLIR", "Execute"]),
                  x + .18, y + 1.30, width - .36, size=11.3)
        if data.get("caption"):
            deck.text(slide, data["caption"], .57, 4.93,
                      8.86, .26, 13, C["focus"], align="center", minimum=12)

    elif layout == "compiler_diversity":
        rows = data.get("rows", [])
        for index, row in enumerate(rows[:2]):
            y = 1.60 + index * 1.47
            deck.rect(slide, .53, y, 8.94, 1.23, C["pale"])
            deck.text(slide, row["title"].replace("Compiler ", "Compiler\n"), .73, y + .35,
                      1.13, .54, 15.5, C["focus"], True, minimum=14)
            xs, widths = (2.03, 4.03, 6.11, 8.07), (1.42, 1.50, 1.40, 1.16)
            for n, (x, width) in enumerate(zip(xs, widths)):
                if n == 1:
                    # The dialect stack is a different plug-in IR world.
                    deck.rect(slide, x, y + .17, width, .88,
                              C["lightgreen"], C["green"])
                    labels = row.get("dialects", [])
                    line_h = .31
                    top = y + .17 + (.88 - len(labels) * line_h) / 2
                    for line, label in enumerate(labels[:2]):
                        deck.text(slide, label, x + .08, top + line * line_h,
                                  width - .16, .27, 14, C["focus"],
                                  mono=True, align="center", minimum=12.5)
                else:
                    deck.rect(slide, x, y + .35, width, .53,
                              C["white"], C["line"])
                    label = row[{0: "frontend", 2: "passes", 3: "target"}[n]]
                    deck.text(slide, label, x + .045, y + .495,
                              width - .09, .29, 14, C["ink"],
                              align="center", minimum=13)
                if n < 3:
                    deck.arrow(slide, x + width + .05, y + .615,
                               xs[n + 1] - .06, y + .615)
        if data.get("caption"):
            deck.text(slide, data["caption"], .55, 4.43,
                      8.90, .35, 12.5, C["focus"], align="center", minimum=11.5)

    elif layout == "controlflow_diversity":
        panels = data.get("panels", [])
        for index, panel in enumerate(panels[:2]):
            x = .53 + index * 4.61
            deck.text(slide, panel["title"], x + .10, 1.52,
                      4.13, .35, 18, C["focus"], True)
            if panel.get("kind") == "tile":
                # A break nested inside the region exits the enclosing loop.
                deck.rect(slide, x + .10, 2.09, 3.36, 1.40,
                          C["pale"], C["line"])
                deck.text(slide, panel["loop_label"], x + .26, 2.22,
                          2.38, .28, 15, C["focus"], True, mono=True)
                deck.rect(slide, x + .40, 2.70, 1.38, .38,
                          C["white"], C["line"])
                deck.text(slide, panel["body_label"], x + .43, 2.795,
                          1.32, .25, 12.5, C["ink"], mono=True, align="center")
                deck.arrow(slide, x + 1.82, 2.89, x + 2.14, 2.89)
                deck.text(slide, "true", x + 1.79, 2.60,
                          .44, .20, 10.2, C["gray"], align="center")
                deck.rect(slide, x + 2.18, 2.67, 1.03, .44,
                          C["lightgreen"], C["green"])
                deck.text(slide, panel["exit_label"], x + 2.21, 2.785,
                          .97, .25, 13, C["focus"], True, mono=True,
                          align="center", minimum=11.5)
                deck.arrow(slide, x + 3.25, 2.89, x + 3.61, 2.89)
                deck.text(slide, "exit\nloop", x + 3.68, 2.70,
                          .57, .43, 10.8, C["gray"], align="center")
                deck.arrow(slide, x + 1.09, 3.10, x + 1.09, 3.30, head=False)
                deck.arrow(slide, x + 1.09, 3.30, x + .24, 3.30, head=False)
                deck.arrow(slide, x + .24, 3.30, x + .24, 2.89, head=False)
                deck.arrow(slide, x + .24, 2.89, x + .37, 2.89)
                deck.text(slide, "false", x + 1.18, 3.15,
                          .66, .20, 10.2, C["gray"])
            else:
                # scf.for has a fixed structured yield/backedge, not break.
                deck.rect(slide, x + .10, 2.09, 4.13, 1.40,
                          C["pale"], C["line"])
                deck.text(slide, panel["loop_label"], x + .26, 2.22,
                          3.77, .28, 15, C["focus"], True, mono=True)
                for nx, width, label in ((x + .43, 1.18, panel["body_label"]),
                                         (x + 2.29, 1.59, panel["exit_label"])):
                    deck.rect(slide, nx, 2.70, width, .38,
                              C["white"], C["line"])
                    deck.text(slide, label, nx + .035, 2.795,
                              width - .07, .24, 12.5, C["ink"],
                              mono=True, align="center", minimum=11.5)
                deck.arrow(slide, x + 1.65, 2.89, x + 2.24, 2.89)
                deck.arrow(slide, x + 3.09, 3.10, x + 3.09, 3.30, head=False)
                deck.arrow(slide, x + 3.09, 3.30, x + 1.02, 3.30, head=False)
                deck.arrow(slide, x + 1.02, 3.30, x + 1.02, 3.12)
                deck.text(slide, panel.get("note", "No direct break in scf.for"),
                          x + .10, 3.70, 4.13, .28, 12, C["gray"], align="center")
        alternative = data.get("alternative")
        if alternative:
            deck.rect(slide, .53, 4.20, 8.94, .43, C["lightgreen"])
            deck.text(slide, alternative["label"], .70, 4.30,
                      3.32, .24, 12, C["focus"], True)
            deck.text(slide, alternative["condition"], 4.16, 4.30,
                      1.50, .25, 12.5, C["focus"], mono=True, align="center")
            deck.arrow(slide, 5.87, 4.405, 6.34, 4.405)
            deck.text(slide, alternative["body"], 6.57, 4.30,
                      2.70, .25, 12.5, C["focus"])

    elif layout == "plugin_architecture":
        # Two user-facing namespaces, a pluggable middle layer, one portable
        # foundation. The common rails show composition without a wiring maze.
        subdsls = data.get("subdsls", [
            dict(title="MlirDSL", body="Reference DSL"),
            dict(title="Your DSL", body="Your dialects"),
        ])
        for index, item in enumerate(subdsls[:2]):
            x = 1.06 + index * 4.65
            deck.rect(slide, x, 1.48, 3.23, .75,
                      C["lightgreen"], C["green"])
            deck.text(slide, item["title"], x + .12, 1.59,
                      2.99, .32, 19, C["focus"], align="center")
            deck.text(slide, item.get("body", ""), x + .12, 1.99,
                      2.99, .20, 10.8, C["gray"], align="center")
        families = data.get("families", [
            dict(title="Dialect", body="Types + operations"),
            dict(title="AST", body="Syntax capture"),
            dict(title="Interop", body="Adapters + exports"),
        ])
        centers = [1.95, 5.00, 8.05]
        deck.arrow(slide, centers[0], 2.48, centers[-1], 2.48, head=False)
        for center in (2.675, 7.325):
            deck.arrow(slide, center, 2.48, center, 2.27)
        for index, item in enumerate(families[:3]):
            x, center = .53 + index * 3.05, centers[index]
            deck.arrow(slide, center, 2.69, center, 2.48, head=False)
            deck.rect(slide, x, 2.72, 2.84, .69, C["white"], C["green"])
            deck.text(slide, item["title"], x + .12, 2.82,
                      2.60, .29, 16, C["focus"], True, align="center")
            if item.get("body"):
                deck.text(slide, item["body"], x + .12, 3.15,
                          2.60, .21, 10.5, C["gray"], align="center")
        deck.arrow(slide, centers[0], 3.64, centers[-1], 3.64, head=False)
        for center in centers:
            deck.arrow(slide, center, 3.64, center, 3.44)
        deck.arrow(slide, 5.00, 3.84, 5.00, 3.64, head=False)
        deck.rect(slide, .53, 3.88, 8.94, .73, C["pale"], C["line"])
        deck.text(slide, data.get("core_label", "BaseDSL"), .73, 4.00,
                  1.60, .34, 18, C["focus"], True)
        labels = data.get("core_labels", ["Types", "Staging", "ABI", "Cache", "Environment", "Diagnostics"])
        left, total, gap = 2.47, 6.79, .075
        width = (total - (len(labels) - 1) * gap) / max(len(labels), 1)
        for index, label in enumerate(labels):
            x = left + index * (width + gap)
            deck.rect(slide, x, 4.08, width, .30, C["white"])
            deck.text(slide, label, x + .025, 4.13, width - .05, .21,
                      10.3, C["ink"], align="center", minimum=9.8)

    elif layout == "plugin_families":
        families = data.get("families", [])
        for index, family in enumerate(families[:3]):
            x, width = .53 + index * 3.05, 2.84
            deck.rect(slide, x, 1.64, width, 2.83, C["pale"])
            deck.rect(slide, x, 1.64, width, .47, C["green"])
            deck.text(slide, family["title"], x + .16, 1.75,
                      width - .32, .31, 17, C["ink"], True)
            deck.text(slide, family.get("protocol", ""), x + .16, 2.32,
                      width - .32, .38, 10.8, C["focus"], mono=True,
                      minimum=10, do_wrap=False)
            for row, role in enumerate(family.get("roles", [])[:2]):
                deck.text(slide, role, x + .16, 2.94 + row * .40,
                          width - .32, .31, 14, C["ink"], minimum=12.5)
            deck.rect(slide, x + .16, 3.92, width - .32, .009, C["line"])
            deck.text(slide, family.get("examples", ""), x + .16, 4.10,
                      width - .32, .24, 10.8, C["focus"], minimum=10)

    elif layout == "plugin_lifecycle":
        stages = data.get("stages", [])
        count, gap = len(stages), .24
        width = (8.94 - gap * (count - 1)) / max(count, 1)
        for index, stage in enumerate(stages):
            x = .53 + index * (width + gap)
            deck.text(slide, f"{index + 1:02}", x, 1.78, width, .24,
                      10.5, C["focus"], True, align="center")
            deck.rect(slide, x, 2.16, width, .54,
                      C["lightgreen"], C["green"])
            deck.text(slide, stage["title"], x + .06, 2.29,
                      width - .12, .30, 14, C["focus"], True,
                      align="center", minimum=12.5)
            if index < count - 1:
                deck.arrow(slide, x + width + .02, 2.43,
                           x + width + gap - .02, 2.43)
            hook_y = 3.03
            for hook in stage.get("hooks", []):
                # Keep long API names legible rather than shrinking all hooks.
                shown = hook.replace("wrap_compiled_function", "wrap_compiled_\nfunction")
                height = .26 * (shown.count("\n") + 1)
                deck.text(slide, shown, x + .035, hook_y,
                          width - .07, height, 11.2, C["ink"],
                          align="center", minimum=10.5, do_wrap=False)
                hook_y += height + .20
            note = stage.get("note", "")
            if note:
                deck.text(slide, note, x + .05, 4.15,
                          width - .10, .47, 9.9, C["gray"],
                          align="center", minimum=9.5)

    elif layout == "emitter_dispatch":
        # The same typed Python expression reaches a dialect-owned emitter.
        # Branch examples are alternatives, not two stages in one lowering.
        nodes = [(.53, 1.28, data.get("expression", "x + y"), 24),
                 (2.19, 1.85, data.get("promotion", "dtype + promotion"), 13.5),
                 (4.43, 2.12, data.get("emitter", "active OpEmitter"), 15.5)]
        for index, (x, width, label, size) in enumerate(nodes):
            deck.rect(slide, x, 2.53, width, .77,
                      C["lightgreen"] if index == 2 else C["pale"],
                      C["green"] if index == 2 else None)
            deck.text(slide, label, x + .06, 2.73,
                      width - .12, .42, size,
                      C["focus"] if index == 2 else C["ink"],
                      align="center", minimum=size - 1.5)
            if index < len(nodes) - 1:
                deck.arrow(slide, x + width + .04, 2.91,
                           nodes[index + 1][0] - .05, 2.91)
        deck.arrow(slide, 6.58, 2.91, 6.83, 2.91, head=False)
        branches = data.get("branches", [])
        for index, branch in enumerate(branches[:2]):
            x, y, width = 7.13, 1.55 + index * 1.60, 2.34
            center = y + .62
            deck.arrow(slide, 6.83, 2.91, 6.83, center, head=False)
            deck.arrow(slide, 6.83, center, x - .05, center)
            deck.rect(slide, x, y, width, 1.24,
                      C["lightgreen"] if index == 1 else C["pale"],
                      C["green"] if index == 1 else C["line"])
            deck.text(slide, branch["title"], x + .12, y + .12,
                      width - .24, .29, 13.5, C["focus"], True,
                      align="center", minimum=12.5)
            deck.text(slide, branch["ir_type"], x + .12, y + .52,
                      width - .24, .29, 13, C["code"], mono=True,
                      align="center", minimum=12)
            deck.text(slide, branch["op"], x + .12, y + .90,
                      width - .24, .23, 11.5, C["gray"], mono=True,
                      align="center", minimum=10.5)
        if data.get("caption"):
            deck.text(slide, data["caption"], .56, 4.14,
                      6.05, .34, 12.5, C["focus"], minimum=11.5)

    elif layout == "ast_plugin_pipeline":
        stages = data.get("stages", [])
        count, gap = len(stages), .25
        width = (8.94 - gap * (count - 1)) / max(count, 1)
        for index, stage in enumerate(stages):
            x = .53 + index * (width + gap)
            deck.rect(slide, x, 1.81, width, 1.44,
                      C["lightgreen"] if index in (1, 3) else C["pale"],
                      C["line"])
            deck.text(slide, stage["title"], x + .13, 2.00,
                      width - .26, .51, 15, C["focus"] if index in (1, 3) else C["ink"],
                      True, minimum=13.5)
            deck.text(slide, stage.get("body", ""), x + .13, 2.76,
                      width - .26, .39, 11.5, C["gray"], minimum=10.5)
            if index < count - 1:
                deck.arrow(slide, x + width + .03, 2.54,
                           x + width + gap - .03, 2.54)
        for owner in data.get("ownership", []):
            start, end = owner["start"], owner["end"]
            x = .53 + start * (width + gap)
            span = (end - start + 1) * width + (end - start) * gap
            deck.rect(slide, x, 3.54, span, .045, C["green"])
            deck.text(slide, owner["label"], x + .035, 3.75,
                      span - .07, .49, 11.8, C["focus"], True,
                      align="center", minimum=10.5)
        targets = data.get("targets", [])
        if targets:
            deck.text(slide, " / ".join(targets), .58, 4.30,
                      8.84, .29, 12, C["gray"], align="center")

    elif layout == "type_rules":
        # A compact reference: available families, staging boundary, conversion
        # rules. Code and MLIR are editable text, not a screenshot.
        families = data.get("families", [])
        family_widths = [1.50, 2.30, 2.50, 2.16]
        x = .53
        for index, family in enumerate(families[:4]):
            width = family_widths[index]
            deck.rect(slide, x, 1.51, width, .38, C["pale"])
            deck.text(slide, family, x + .07, 1.60,
                      width - .14, .23, 11.4, C["gray"],
                      align="center", minimum=10.8)
            x += width + .16
        pipeline = data.get("pipeline", ["Python literal", "Typed wrapper", "MLIR value"])
        for index, label in enumerate(pipeline[:3]):
            x = .83 + index * 3.10
            deck.text(slide, label, x, 2.12, 2.15, .31,
                      16, C["focus"] if index else C["ink"],
                      align="center", minimum=15)
            if index < 2:
                deck.arrow(slide, x + 2.27, 2.25, x + 2.91, 2.25)
        cols = (.69, 2.33, 5.86)
        widths = (1.44, 3.33, 3.40)
        for x, width, label in zip(cols, widths, ("Conversion", "Python", "MLIR")):
            deck.text(slide, label, x, 2.58, width, .25,
                      10.3, C["gray"], True)
        deck.rect(slide, .53, 2.87, 8.94, .012, C["line"])
        rules = data.get("rules", [])
        for index, rule in enumerate(rules[:6]):
            y = 2.94 + index * .292
            deck.text(slide, rule["label"], cols[0], y,
                      widths[0], .28, 12.5, C["focus"], True, minimum=11.5)
            deck.text(slide, rule["example"], cols[1], y,
                      widths[1], .28, 11.5, C["code"], mono=True,
                      do_wrap=False, minimum=10.2)
            deck.text(slide, rule["op"], cols[2], y,
                      widths[2], .28, 11.5, C["code"], mono=True,
                      do_wrap=False, minimum=10.2)

    elif layout == "design_simple":
        for index, item in enumerate(items):
            y = 1.81 + index * 1.10
            deck.text(slide, f"0{index + 1}", .56, y + .03, .47, .40,
                      18, C["focus"], True)
            deck.text(slide, item["title"], 1.28, y, 8.15, .53,
                      23, minimum=20)
            if item.get("body"):
                deck.text(slide, item["body"], 1.30, y + .63,
                          8.10, .42, 13, C["gray"], minimum=12)
        if data.get("takeaway"):
            _statement(deck, slide, data["takeaway"], .55, 4.38,
                       8.90, .39, 16, data.get("bold_phrase", "multi-stage programming"))

    elif layout == "generator_idea":
        # One coherent reading path: Python source -> its execution -> emitted IR.
        deck.text(slide, "Python", .54, 1.53, 3.32, .27,
                  12, C["ink"], True)
        deck.text(slide, "Trace: execute Python", 3.96, 1.53, 2.08, .27,
                  11, C["focus"], True, align="center")
        deck.text(slide, "MLIR", 6.18, 1.53, 3.28, .27,
                  12, C["focus"], True)
        deck.code(slide, data["left"], .50, 1.99, 3.33, 2.03, size=10.6)
        deck.code(slide, data["right"], 6.17, 1.99, 3.33, 2.03,
                  size=10.6, green=True)
        deck.arrow(slide, 3.84, 2.90, 3.99, 2.90, head=False)
        deck.arrow(slide, 3.99, 2.50, 3.99, 3.54, head=False)
        deck.arrow(slide, 3.99, 2.50, 4.14, 2.50, C["gray"])
        deck.arrow(slide, 3.99, 3.54, 4.14, 3.54)
        deck.arrow(slide, 5.88, 3.54, 6.15, 3.54)
        for index, item in enumerate(items[:2]):
            y = 2.12 + index * 1.04
            deck.rect(slide, 4.16, y, 1.69, .76,
                      C["pale"] if index == 0 else C["lightgreen"])
            deck.text(slide, item["title"], 4.25, y + .10, 1.51, .24,
                      12, C["gray"] if index == 0 else C["focus"],
                      True, align="center", minimum=11)
            deck.text(slide, item["body"], 4.25, y + .39, 1.51, .35,
                      10.2, C["gray"], align="center", minimum=9.7)
        deck.text(slide, "Objects · imports · metaprogramming", .54, 4.15,
                  4.00, .27, 10.1, C["gray"])
        targets = data.get("targets", ["Compile", "Test", "Consume IR"])
        # These are alternative consumers of the generated IR, not more passes.
        origin = 7.835
        deck.arrow(slide, origin, 4.03, origin, 4.12, head=False)
        target_w = 3.33 / len(targets)
        for index, target in enumerate(targets):
            x = 6.17 + index * target_w
            center = x + target_w / 2
            deck.arrow(slide, origin, 4.12, center, 4.12, head=False)
            deck.arrow(slide, center, 4.12, center, 4.24)
            deck.text(slide, target, x + .03, 4.29,
                      target_w - .06, .34, 9.9, C["focus"],
                      align="center", minimum=9.3)

    elif layout == "calling":
        node_x = (.73, 3.94, 7.15)
        labels = data.get("nodes", ["Python", "@jit", "@kernel"])
        for index, (x, label) in enumerate(zip(node_x, labels)):
            _box(deck, slide, label, x, 1.83, 2.10, .64,
                 accent=index > 0, size=18)
            if index < 2:
                deck.arrow(slide, x + 2.15, 2.15, node_x[index + 1] - .07, 2.15)
        deck.text(slide, "compile / invoke", 2.91, 1.61, .95, .39,
                  9.3, C["gray"], align="center")
        deck.text(slide, "GPU launch", 6.14, 1.72, .92, .28,
                  9.3, C["gray"], align="center")
        # The feedback line encodes trace-time Python helper execution.
        deck.arrow(slide, 4.99, 2.48, 4.99, 2.81, head=False)
        deck.arrow(slide, 4.99, 2.81, 1.77, 2.81, head=False)
        deck.arrow(slide, 1.77, 2.81, 1.77, 2.51)
        deck.text(slide, "Python helpers run while tracing", 1.99, 2.90,
                  3.51, .28, 10.3, C["focus"], align="center")
        for index, item in enumerate(items[:3]):
            x = .62 + index * 3.21
            deck.text(slide, item["title"], x, 3.56, 2.75, .38,
                      13, C["focus"], True, minimum=12)
            deck.text(slide, item["body"], x, 4.00, 2.75, .54,
                      11.3, C["gray"], minimum=10.5)

    elif layout == "capture_pipeline":
        stages = data.get("stages", [])
        n = len(stages)
        gap = .27
        width = (8.94 - (n - 1) * gap) / n
        for index, stage in enumerate(stages):
            x = .53 + index * (width + gap)
            accent = index in {1, 3}
            deck.rect(slide, x, 1.85, width, 1.58,
                      C["lightgreen"] if accent else C["pale"], C["line"])
            deck.text(slide, stage["title"], x + .13, 2.04,
                      width - .26, .55, 13, C["focus"] if accent else C["ink"],
                      True, minimum=11.5)
            deck.text(slide, stage["body"], x + .13, 2.67,
                      width - .26, .56, 10.4, C["gray"], minimum=9.7)
            if index < n - 1:
                deck.arrow(slide, x + width + .035, 2.64,
                           x + width + gap - .035, 2.64)
        targets = data.get("targets", ["scf", "mydialect.cf"])
        deck.text(slide, "Same capture · pluggable builders", .59, 3.93,
                  4.20, .33, 13, C["focus"], minimum=12)
        last_center = .53 + (n - 1) * (width + gap) + width / 2
        deck.arrow(slide, last_center, 3.45, last_center, 3.71, head=False)
        for index, target in enumerate(targets):
            x = 5.15 + index * 2.21
            deck.arrow(slide, last_center, 3.71, x + .97, 3.71, head=False)
            deck.arrow(slide, x + .97, 3.71, x + .97, 3.90)
            _box(deck, slide, target, x, 3.93, 1.95, .54,
                 accent=index == 1, size=13)

    elif layout == "type_storage":
        _box(deck, slide, data.get("root_label", "Numeric.value"),
             3.73, 1.63, 2.54, .62, accent=True, size=17)
        for cx in (2.67, 7.31):
            deck.arrow(slide, 5.00, 2.28, 5.00, 2.45, head=False)
            deck.arrow(slide, 5.00, 2.45, cx, 2.45, head=False)
            deck.arrow(slide, cx, 2.45, cx, 2.68)
        if data.get("left") and data.get("right"):
            deck.code(slide, data["left"], .51, 2.75, 4.30, 1.74, size=11)
            deck.code(slide, data["right"], 5.19, 2.75, 4.30, 1.74,
                      size=11, green=True)
        else:
            for index, item in enumerate(items[:2]):
                x = .66 + 4.67 * index
                deck.text(slide, item["title"], x, 2.83, 4.03, .51,
                          18, C["focus"], True)
                deck.text(slide, item["body"], x, 3.51, 4.03, .74,
                          13, C["gray"], minimum=12)
    else:
        return False
    return True
